"""批量查找替换模块:在 docx / xlsx / pptx 中替换文本,支持普通文本与正则表达式。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from pptx import Presentation


def _build_replacer(pairs: list[tuple[str, str]],
                    use_regex: bool) -> Callable[[str], tuple[str, int]]:
    """根据规则构建替换函数:text -> (新文本, 命中次数)。

    use_regex 为 True 时,查找内容按正则表达式解释,替换为中支持
    分组引用(如 \\1);无效正则在此处即抛出 ValueError。
    """
    if use_regex:
        compiled: list[tuple[re.Pattern, str]] = []
        for old, new in pairs:
            if not old:
                continue
            try:
                compiled.append((re.compile(old), new))
            except re.error as exc:
                raise ValueError(f'无效的正则表达式「{old}」: {exc}') from exc
        if not compiled:
            raise ValueError('正则模式下查找内容不能为空')

        def apply_regex(text: str) -> tuple[str, int]:
            if not text:  # 空文本跳过,避免 x* 类可匹配空串的正则误插入内容
                return text, 0
            total = 0
            for pattern, new in compiled:
                text, n = pattern.subn(new, text)
                total += n
            return text, total

        apply_regex.is_regex = True  # 标记:段落级替换必须按整段文本匹配
        return apply_regex

    def apply_plain(text: str) -> tuple[str, int]:
        total = 0
        for old, new in pairs:
            if old and old in text:
                total += text.count(old)
                text = text.replace(old, new)
        return text, total

    apply_plain.is_regex = False
    return apply_plain


def _replace_docx_paragraph(paragraph, apply: Callable[[str], tuple[str, int]]) -> int:
    """替换段落中的文本。

    普通文本:优先逐 run 替换以保留格式;关键字跨 run 断裂时,合并
    整段文本替换后写回第一个 run(格式退化为首 run 样式)。
    正则模式:锚点(^ $ \\b)在单个 run 上的语义与整段不同,因此
    跳过逐 run 分支,只按整段文本匹配,保证结果正确。
    """
    # 正则模式不做逐 run 替换,避免锚点语义漂移
    if not getattr(apply, 'is_regex', False):
        # 先尝试逐 run 替换（不破坏格式）
        inline_count = 0
        for run in paragraph.runs:
            new_text, n = apply(run.text)
            if n:
                run.text = new_text
                inline_count += n
        if inline_count:
            return inline_count

    # 按整段文本匹配(普通文本的关键字可能跨 run;正则始终走这里)
    replaced, count = apply(paragraph.text)
    if count == 0:
        return 0
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(replaced)
    return count


def _replace_docx(path: Path, apply) -> int:
    doc = Document(str(path))
    total = 0

    def walk_paragraphs(paragraphs):
        nonlocal total
        for p in paragraphs:
            total += _replace_docx_paragraph(p, apply)

    walk_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk_paragraphs(cell.paragraphs)
    # 页眉页脚
    for section in doc.sections:
        for hf in (section.header, section.footer):
            walk_paragraphs(hf.paragraphs)
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        walk_paragraphs(cell.paragraphs)
    doc.save(str(path))
    return total


def _replace_xlsx(path: Path, apply) -> int:
    wb = load_workbook(str(path))
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                    continue
                new_val, n = apply(cell.value)
                if n:
                    total += n
                    cell.value = new_val
    wb.save(str(path))
    return total


def _replace_pptx_textframe(tf, apply) -> int:
    count = 0
    use_per_run = not getattr(apply, 'is_regex', False)  # 正则模式只按整段匹配
    for para in tf.paragraphs:
        # 普通文本:先逐 run 替换以保留格式
        inline = 0
        if use_per_run:
            for run in para.runs:
                new_text, n = apply(run.text)
                if n:
                    run.text = new_text
                    inline += n
        if inline:
            count += inline
            continue
        # 整段文本匹配(跨 run 关键字 / 正则模式)
        replaced, hit = apply(''.join(run.text for run in para.runs))
        if hit:
            count += hit
            if para.runs:
                para.runs[0].text = replaced
                for run in para.runs[1:]:
                    run.text = ''
            else:
                para.text = replaced
    return count


def _replace_pptx(path: Path, apply) -> int:
    prs = Presentation(str(path))
    total = 0

    def walk_shapes(shapes):
        nonlocal total
        for shape in shapes:
            if shape.has_text_frame:
                total += _replace_pptx_textframe(shape.text_frame, apply)
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        total += _replace_pptx_textframe(cell.text_frame, apply)
            if shape.shape_type == 6:
                walk_shapes(shape.shapes)

    for slide in prs.slides:
        walk_shapes(slide.shapes)
    prs.save(str(path))
    return total


def replace_in_file(path: str | Path, pairs: list[tuple[str, str]],
                    use_regex: bool = False) -> int:
    """在单个文件中执行替换,返回替换发生的次数(直接保存原文件)。

    use_regex 为 True 时查找内容按正则表达式解释(替换为中可用 \\1 等分组引用)。
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f'文件不存在: {path.name}')
    if not pairs:
        return 0
    apply = _build_replacer(pairs, use_regex)
    ext = path.suffix.lower()
    if ext == '.docx':
        return _replace_docx(path, apply)
    if ext in ('.xlsx', '.xlsm'):
        return _replace_xlsx(path, apply)
    if ext == '.pptx':
        return _replace_pptx(path, apply)
    raise ValueError(f'不支持的文件类型: {path.name}(仅支持 docx/xlsx/pptx)')
