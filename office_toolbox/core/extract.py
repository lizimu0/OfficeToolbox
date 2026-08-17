"""文本与图片提取模块:从 docx / xlsx / pptx 中提取文字和图片。"""
from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation


def _extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    lines: list[str] = []
    for block in doc.element.body:
        # 按文档顺序处理段落与表格
        tag = block.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            text = Paragraph(block, doc).text.strip()
            if text:
                lines.append(text)
        elif tag == 'tbl':
            table = Table(block, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(' | '.join(cells))
            lines.append('')
    return '\n'.join(lines)


def _extract_xlsx_text(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f'== 工作表: {ws.title} ==')
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            parts.append('\t'.join('' if v is None else str(v) for v in row))
        parts.append('')
    wb.close()
    return '\n'.join(parts)


def _pptx_shape_text(shape, lines: list[str]) -> None:
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = ''.join(run.text for run in para.runs).strip()
            if text:
                lines.append(text)
    if getattr(shape, 'has_table', False) and shape.has_table:
        for row in shape.table.rows:
            lines.append(' | '.join(cell.text.strip() for cell in row.cells))
    if shape.shape_type == 6:  # 组合图形,递归提取
        for sub in shape.shapes:
            _pptx_shape_text(sub, lines)


def _extract_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    lines: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        lines.append(f'== 第 {idx} 页幻灯片 ==')
        for shape in slide.shapes:
            _pptx_shape_text(shape, lines)
        lines.append('')
    return '\n'.join(lines)


def extract_text(path: str | Path) -> str:
    """根据扩展名自动分发,返回提取出的纯文本。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f'文件不存在: {path.name}')
    ext = path.suffix.lower()
    if ext == '.docx':
        return _extract_docx_text(path)
    if ext in ('.xlsx', '.xlsm'):
        return _extract_xlsx_text(path)
    if ext == '.pptx':
        return _extract_pptx_text(path)
    raise ValueError(f'不支持的文件类型: {path.name}(仅支持 docx/xlsx/pptx)')


_IMAGE_DIRS = {'.docx': 'word/media/', '.xlsx': 'xl/media/', '.pptx': 'ppt/media/'}


def extract_images(path: str | Path, out_dir: str | Path) -> list[str]:
    """从 Office 文件中导出所有嵌入图片,返回保存的文件名列表。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f'文件不存在: {path.name}')
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    prefix = _IMAGE_DIRS.get(path.suffix.lower())
    if prefix is None:
        raise ValueError(f'不支持的文件类型: {path.name}')
    saved: list[str] = []
    try:
        with zipfile.ZipFile(str(path)) as zf:
            for name in zf.namelist():
                if name.startswith(prefix) and not name.endswith('/'):
                    data = zf.read(name)
                    out_name = path.stem + '_' + Path(name).name
                    (out_dir / out_name).write_bytes(data)
                    saved.append(out_name)
    except zipfile.BadZipFile as exc:
        raise ValueError(f'无法读取「{path.name}」: 文件已损坏或不是有效的 Office 文档') from exc
    return saved
