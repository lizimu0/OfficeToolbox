"""批量重命名模块:文本替换、前缀序号、Excel 映射三种规则。"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


def _unique_name(target: Path) -> Path:
    """目标名已被占用时自动追加序号。"""
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    n = 2
    while True:
        candidate = target.with_name(f'{stem}_{n}{suffix}')
        if not candidate.exists():
            return candidate
        n += 1


def _safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', '_', name).strip()


def rename_by_replace(paths: list[str | Path], old: str, new: str,
                      preview: bool = False) -> list[tuple[str, str]]:
    """文件名中的文本替换。返回 [(旧名, 新名)] 仅包含实际改名的项。

    preview=True 时只计算结果不实际改名(用于界面预演)。
    """
    results: list[tuple[str, str]] = []
    for p in paths:
        src = Path(p)
        if old and old in src.stem:
            target = src.with_name(_safe_name(src.stem.replace(old, new)) + src.suffix)
            target = _unique_name(target)
            if not preview:
                src.rename(target)
            results.append((src.name, target.name))
    return results


def rename_by_sequence(paths: list[str | Path], prefix: str,
                       start: int = 1, width: int = 3,
                       preview: bool = False) -> list[tuple[str, str]]:
    """按列表顺序重命名为 前缀+序号。返回 [(旧名, 新名)]。

    preview=True 时只计算结果不实际改名。
    """
    results: list[tuple[str, str]] = []
    for i, p in enumerate(paths, start):
        src = Path(p)
        target = src.with_name(f'{_safe_name(prefix)}{i:0{width}d}{src.suffix}')
        target = _unique_name(target)
        if not preview:
            src.rename(target)
        results.append((src.name, target.name))
    return results


def _load_mapping(xlsx_path: str | Path) -> dict[str, str]:
    """读取映射表:首列为原文件名(可不含扩展名),次列为新文件名。"""
    wb = load_workbook(str(xlsx_path), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    mapping: dict[str, str] = {}
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0 or row[0] is None:
            continue
        old = str(row[0]).strip()
        new = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
        if old and new:
            mapping[Path(old).stem.lower()] = _safe_name(new)
    wb.close()
    return mapping


def rename_by_mapping(paths: list[str | Path],
                      xlsx_path: str | Path,
                      preview: bool = False) -> tuple[list[tuple[str, str]], list[str]]:
    """按 Excel 映射表重命名。返回 (已改名列表, 未匹配文件列表)。

    preview=True 时只计算结果不实际改名。
    """
    mapping = _load_mapping(xlsx_path)
    if not mapping:
        raise ValueError('映射表中没有有效的映射数据')
    renamed: list[tuple[str, str]] = []
    skipped: list[str] = []
    for p in paths:
        src = Path(p)
        new_stem = mapping.get(src.stem.lower())
        if new_stem is None:
            skipped.append(src.name)
            continue
        target = src.with_name(new_stem + src.suffix)
        target = _unique_name(target)
        if not preview:
            src.rename(target)
        renamed.append((src.name, target.name))
    return renamed, skipped
