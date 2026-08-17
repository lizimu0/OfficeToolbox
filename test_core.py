# -*- coding: utf-8 -*-
"""核心功能冒烟测试:生成样例文件并逐一验证各模块。"""
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from office_toolbox.core import extract, merge_excel, pdf_tools, rename, replace, template
from office_toolbox.core.convert import csv_to_xlsx

tmp = Path(tempfile.mkdtemp(prefix='office_tool_test_'))
print(f'测试目录: {tmp}')


def check(name, cond):
    print(('  [通过] ' if cond else '  [失败] ') + name)
    if not cond:
        sys.exit(1)


# ---- 生成样例文件 ----
doc = Document()
doc.add_paragraph('甲方:某某科技有限公司')
doc.add_paragraph('合同金额为人民币100万元。')
table = doc.add_table(rows=1, cols=2)
table.cell(0, 0).text = '项目'
table.cell(0, 1).text = '备注'
docx_path = tmp / '样例.docx'
doc.save(str(docx_path))

wb = Workbook()
ws = wb.active
ws.append(['姓名', '部门', '工资'])
ws.append(['张三', '研发部', 12000])
ws.append(['李四', '市场部', 9000])
xlsx_path = tmp / '样例.xlsx'
wb.save(str(xlsx_path))

wb2 = Workbook()
ws2 = wb2.active
ws2.append(['姓名', '部门', '工资'])
ws2.append(['王五', '财务部', 10000])
xlsx_path2 = tmp / '样例2.xlsx'
wb2.save(str(xlsx_path2))

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = '季度汇报'
slide.placeholders[1].text = '某某科技有限公司业绩'
pptx_path = tmp / '样例.pptx'
prs.save(str(pptx_path))

# ---- 1. 文本提取 ----
print('1. 文本提取')
check('docx 提取包含段落与表格', '某某科技有限公司' in extract.extract_text(docx_path)
      and '备注' in extract.extract_text(docx_path))
check('xlsx 提取包含表头与数据', '张三' in extract.extract_text(xlsx_path)
      and '部门' in extract.extract_text(xlsx_path))
check('pptx 提取包含标题与正文', '季度汇报' in extract.extract_text(pptx_path))

# ---- 2. 批量替换 ----
print('2. 批量替换')
n = replace.replace_in_file(docx_path, [('某某科技有限公司', '星河集团')])
check(f'docx 替换 {n} 处', n >= 1 and '星河集团' in extract.extract_text(docx_path))
n = replace.replace_in_file(xlsx_path, [('研发部', '研发中心')])
check(f'xlsx 替换 {n} 处', n == 1 and '研发中心' in extract.extract_text(xlsx_path))
n = replace.replace_in_file(pptx_path, [('某某科技', '星河')])
check(f'pptx 替换 {n} 处', n >= 1 and '星河' in extract.extract_text(pptx_path))

# 正则替换(含分组引用)
n = replace.replace_in_file(docx_path, [(r'(\d+)万元', r'\1万元整')], use_regex=True)
regex_text = extract.extract_text(docx_path)
check(f'docx 正则替换 {n} 处', n >= 1 and '100万元整' in regex_text)
try:
    replace.replace_in_file(docx_path, [('([unclosed', 'x')], use_regex=True)
    check('无效正则报错', False)
except ValueError:
    check('无效正则报错', True)

# 正则锚点语义:^ $ 按“整段”解释,不在 run 边界误生效
anchor_doc = Document()
ap = anchor_doc.add_paragraph()
ap.add_run('foo')
ap.add_run('bar')
anchor_path = tmp / '锚点.docx'
anchor_doc.save(str(anchor_path))
n = replace.replace_in_file(anchor_path, [('^foo$', 'X')], use_regex=True)
check('锚点正则不误匹配单 run', n == 0 and 'foobar' in extract.extract_text(anchor_path))
n = replace.replace_in_file(anchor_path, [('foobar', 'X')], use_regex=True)
check('正则整段命中跨 run 文本', n == 1 and extract.extract_text(anchor_path).strip() == 'X')

# ---- 3. Excel 合并 ----
print('3. Excel 合并')
out_sheets = tmp / '合并_按表.xlsx'
count = merge_excel.merge_as_sheets([xlsx_path, xlsx_path2], out_sheets)
check(f'按工作表合并({count} 个表)', count == 2 and out_sheets.exists())
out_rows = tmp / '合并_按行.xlsx'
rows = merge_excel.merge_append_rows([xlsx_path, xlsx_path2], out_rows)
text = extract.extract_text(out_rows)
check(f'按行追加合并({rows} 行)', rows == 3 and '张三' in text and '王五' in text)

# ---- 4. 模板批量生成 ----
print('4. 模板批量生成')
tpl = Document()
tpl.add_paragraph('聘用通知:{{姓名}}')
tpl.add_paragraph('部门:{{部门}},薪资:{{工资}}')
tpl_path = tmp / '模板.docx'
tpl.save(str(tpl_path))
out_dir = tmp / '生成结果'
files = template.generate_from_template(tpl_path, xlsx_path, out_dir, name_field='姓名')
check(f'生成 {len(files)} 个文档', len(files) == 2 and files == ['张三.docx', '李四.docx'])
check('文档内容已填充', '张三' in extract.extract_text(out_dir / '张三.docx')
      and '{{' not in extract.extract_text(out_dir / '张三.docx'))

# 图片占位符:数据源“照片”列为图片路径,生成文档应内嵌图片
from PIL import Image  # noqa: E402

img_path = tmp / '照片.png'
Image.new('RGB', (60, 40), (200, 60, 60)).save(img_path)
wb_photo = Workbook()
ws_photo = wb_photo.active
ws_photo.append(['姓名', '照片'])
ws_photo.append(['王五', str(img_path)])
photo_xlsx = tmp / '照片数据.xlsx'
wb_photo.save(photo_xlsx)
tpl2 = Document()
tpl2.add_paragraph('姓名:{{姓名}}')
bold_p = tpl2.add_paragraph()
bold_run = bold_p.add_run('照片:{{图片:照片}}')
bold_run.bold = True  # 验证图片段落重建后格式保留
tpl2_path = tmp / '带图模板.docx'
tpl2.save(str(tpl2_path))
photo_out = tmp / '带图结果'
files2 = template.generate_from_template(tpl2_path, photo_xlsx, photo_out)
check(f'带图模板生成 {len(files2)} 个文档', len(files2) == 1)
with zipfile.ZipFile(photo_out / '文档1.docx') as zf:
    media = [n for n in zf.namelist() if n.startswith('word/media/')]
check(f'文档内嵌图片({len(media)} 张)', len(media) == 1)
photo_text = extract.extract_text(photo_out / '文档1.docx')
check('图片占位符已清除且文字保留', '姓名:王五' in photo_text and '{{' not in photo_text)

# 图片段落重建后应继承原段首 run 的字符格式(此处验证加粗)
out_doc = Document(str(photo_out / '文档1.docx'))
img_para = next(p for p in out_doc.paragraphs if '照片' in p.text)
text_runs = [r for r in img_para.runs if r.text]
check('图片段落文字保留原格式(加粗)', bool(text_runs) and all(r.bold for r in text_runs))

# ---- 5. CSV → XLSX ----
print('5. CSV 转 XLSX')
csv_path = tmp / '样例.csv'
csv_path.write_text('编号,名称\n1,测试项\n', encoding='utf-8')
xlsx_out = tmp / 'csv转换.xlsx'
csv_to_xlsx(csv_path, xlsx_out)
check('csv 转 xlsx', '测试项' in extract.extract_text(xlsx_out))

# ---- 6. PDF 合并拆分与 Word 合并 ----
print('6. PDF 合并拆分 / Word 合并')
from pypdf import PdfReader, PdfWriter  # noqa: E402


def make_pdf(path, pages):
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=300)
    with open(path, 'wb') as f:
        w.write(f)


pdf1, pdf2 = tmp / 'a.pdf', tmp / 'b.pdf'
make_pdf(pdf1, 2)
make_pdf(pdf2, 3)
merged_pdf = tmp / '合并.pdf'
n = pdf_tools.merge_pdfs([pdf1, pdf2], merged_pdf)
check(f'PDF 合并({n} 页)', n == 5 and len(PdfReader(str(merged_pdf)).pages) == 5)
split_out = tmp / '拆分结果'
files = pdf_tools.split_pdf(merged_pdf, split_out, pages_per_file=2)
check(f'PDF 拆分(每 2 页一份, {len(files)} 个)', len(files) == 3
      and all((split_out / f).exists() for f in files))

d1 = Document()
d1.add_paragraph('第一篇内容')
d1_path = tmp / 'w1.docx'
d1.save(str(d1_path))
d2 = Document()
d2.add_paragraph('第二篇内容')
d2_path = tmp / 'w2.docx'
d2.save(str(d2_path))
word_out = tmp / '合并文档.docx'
pdf_tools.merge_word([d1_path, d2_path], word_out)
text = extract.extract_text(word_out)
check('Word 合并', '第一篇内容' in text and '第二篇内容' in text)

# ---- 7. 批量重命名 ----
print('7. 批量重命名')
rn_dir = tmp / '重命名'
rn_dir.mkdir()
f1 = rn_dir / '旧名称_报告.docx'
f2 = rn_dir / '旧名称_合同.docx'
f1.write_bytes(Path(d1_path).read_bytes())
f2.write_bytes(Path(d2_path).read_bytes())
res = rename.rename_by_replace([f1, f2], '旧名称', '新名称')
check('文件名文本替换', len(res) == 2 and (rn_dir / '新名称_报告.docx').exists())
res = rename.rename_by_sequence(sorted(rn_dir.glob('*.docx')), '文档_', start=1)
check('前缀序号重命名', (rn_dir / '文档_001.docx').exists() and (rn_dir / '文档_002.docx').exists())
map_wb = Workbook()
map_ws = map_wb.active
map_ws.append(['原文件名', '新文件名'])
map_ws.append(['文档_001', '第一季度报告'])
map_path = tmp / '映射表.xlsx'
map_wb.save(str(map_path))
renamed, skipped = rename.rename_by_mapping(sorted(rn_dir.glob('*.docx')), map_path)
check('Excel 映射重命名', len(renamed) == 1 and len(skipped) == 1
      and (rn_dir / '第一季度报告.docx').exists())

# ---- 8. GUI 可导入 ----
print('8. GUI 导入检查')
from office_toolbox.gui.main_window import MainWindow  # noqa: E402
check('主窗口模块导入成功', MainWindow is not None)

print('\n全部测试通过')
